import hashlib
import io
import json
import os
import tempfile
import zipfile
from pathlib import Path
import pytest
from ka_detector.documents import extract
from ka_detector.registry import model_manifest
from ka_detector.research import interval, load_corpus, metrics

TEXT = (
    "В мастерской заменили старый кабель. После проверки оказалось, что разъём плохо держался в корпусе. Мы записали результат и подготовили запасную деталь. "
    * 5
)


@pytest.mark.parametrize(
    "name,data",
    [
        ("pdf", b"not a pdf"),
        ("docx", b"not zip"),
        ("txt", b"\0" * 200),
        ("exe", b"hello" * 40),
    ],
)
def test_bad_documents(name, data):
    with pytest.raises((ValueError, UnicodeError)):
        extract(data, name)


def test_txt_preserves_paragraphs_and_emoji():
    value = TEXT + "\n\nНовый абзац. 🙂"
    assert extract(value.encode(), "txt")["text"] == value


def test_docx_and_zip_expansion():
    from docx import Document

    document = Document()
    document.add_paragraph(TEXT)
    document.add_paragraph("Второй абзац.")
    buffer = io.BytesIO()
    document.save(buffer)
    assert "\n\nВторой абзац." in extract(buffer.getvalue(), "docx")["text"]
    bad = io.BytesIO()
    with zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "a" * 100000)
    with pytest.raises(ValueError):
        extract(bad.getvalue(), "docx")


def test_scanned_pdf_is_not_ocr():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match="no_text_layer"):
        extract(buffer.getvalue(), "pdf")


def test_registry_rejects_path_and_corruption(tmp_path, monkeypatch):
    from ka_detector import registry

    monkeypatch.setattr(registry, "ROOT", tmp_path)
    with pytest.raises(ValueError):
        model_manifest("../secret")
    d = tmp_path / "ka-test"
    d.mkdir()
    (d / "weight").write_bytes(b"original")
    (d / "manifest.json").write_text(
        json.dumps(
            {
                "id": "ka-test",
                "threshold": 0.9,
                "status": "local-research",
                "sha256": {"weight": hashlib.sha256(b"original").hexdigest()},
            }
        )
    )
    model_manifest("ka-test")
    (d / "weight").write_bytes(b"corrupt")
    with pytest.raises(ValueError, match="model_corrupt"):
        model_manifest("ka-test")


def test_intervals_do_not_claim_three_percent():
    assert interval(0, 12)[1] > 0.2
    assert metrics([]) == {"n": 0}


def test_smoke_cannot_be_registered(tmp_path):
    from ka_detector.register import register

    (tmp_path / "run.json").write_text(json.dumps({"kind": "synthetic-gradient-smoke"}))
    with pytest.raises(ValueError, match="Smoke artifacts"):
        register(tmp_path, "ka-not-a-real-model")


def test_corpus_rejects_source_leakage(tmp_path):
    common = {
        "genre": "news",
        "language": "ru",
        "source": "test",
        "sourceGroup": "same-document",
        "license": "CC0",
        "provenance": "synthetic QA fixture",
        "generator": "test",
        "edited": False,
        "addedAt": "2026-10-27",
    }
    rows = [
        dict(
            common,
            id="a",
            text="Разные слова для первого примера.",
            label="human",
            split="train",
        ),
        dict(
            common,
            id="b",
            text="Совсем иной текст второго примера.",
            label="ai",
            split="test",
        ),
    ]
    path = tmp_path / "corpus.jsonl"
    path.write_text("\n".join(json.dumps(row) for row in rows))
    with pytest.raises(ValueError, match="Source-group leakage"):
        load_corpus(path)


def test_real_v1_is_not_a_placeholder():
    from ka_detector.engine import analyze

    result = analyze(TEXT, ["upstream-v1"])["results"][0]
    assert result["aggregation"] == "upstream_v1_features"
    assert 0 <= result["score"] <= 1


def test_upstream_probability_compatibility(tmp_path):
    from aidetector import transformer_score as upstream
    from ka_detector.engine import analyze

    m = model_manifest("upstream-v2")
    upstream.REGISTRY["ka-test-v2"] = {
        "model": str(m["path"] / "v2_model"),
        "calib": str(m["path"] / "v2_calibrator.joblib"),
        "thr": str(m["path"] / "threshold_v2.json"),
    }
    scorer = upstream.Scorer("ka-test-v2")
    scorer._cache_path = str(tmp_path / "cache.json")
    original = float(scorer.proba([TEXT])[0])
    actual = analyze(TEXT, ["upstream-v2"])["results"][0]
    assert actual["windows"] == 1
    assert abs(original - actual["score"]) < 1e-5


def test_real_engine_map_and_short_text():
    from ka_detector.engine import analyze

    result = analyze(TEXT, ["upstream-v2"])
    assert result["words"] > 40 and len(result["results"]) == 1
    assert 0 <= result["results"][0]["score"] <= 1
    p = result["results"][0]["paragraphs"][0]
    assert TEXT[p["start"] : p["end"]] == TEXT.strip()
    short = analyze(
        "Короткая фраза о работе программы и настройках компьютера.", ["upstream-v2"]
    )
    assert short["results"][0]["verdict"] == "insufficient"


def test_internal_api_permissions_and_limits():
    import httpx

    base = "http://ka-text-detector:8020"
    with httpx.Client(base_url=base, timeout=90) as client:
        assert client.get("/health").status_code == 200
        assert client.get("/v1/status").status_code == 401
        client.headers["authorization"] = (
            "Bearer " + os.environ["DETECTOR_SERVICE_TOKEN"]
        )
        assert (
            client.post(
                "/v1/analyze", json={"text": "a", "models": ["upstream-v2"]}
            ).status_code
            == 422
        )
        assert (
            client.post(
                "/v1/analyze", json={"text": TEXT, "models": ["unknown"]}
            ).status_code
            == 409
        )
        response = client.post(
            "/v1/analyze", json={"text": TEXT, "models": ["upstream-v2"]}
        )
        assert response.status_code == 200, response.text
        assert response.json()["data"]["results"][0]["model"] == "upstream-v2"
